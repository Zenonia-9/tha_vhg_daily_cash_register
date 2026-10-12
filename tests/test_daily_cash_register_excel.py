# -*- coding: utf-8 -*-

from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock
from zipfile import ZipFile
from lxml import etree

from odoo.tests.common import TransactionCase

from odoo.addons.tha_vhg_daily_cash_register.models.daily_cash_register_xlsx import (
    COLOR_BALANCE,
    COLOR_HEADER,
    COLOR_TOTAL,
)


def _fill_rgb(cell):
    fill = cell.fill
    color = getattr(fill.fgColor, "rgb", None) or getattr(fill.start_color, "rgb", None)
    if not color:
        return ""
    return str(color).upper()[-6:]


class TestDailyCashRegisterExcel(TransactionCase):
    def setUp(self):
        super().setUp()
        self.company = self.env.company
        self.register = self.env["vhg.daily.cash.register"].new(
            {
                "name": "DCR/TEST",
                "date": date(2026, 4, 4),
                "company_id": self.company.id,
                "opening_kyats": 1029986000.0,
                "opening_sgd": 0.0,
                "opening_baht": 0.0,
                "opening_usd": 0.0,
                "line_ids": [
                    (
                        0,
                        0,
                        {
                            "line_type": "receipt",
                            "sequence": 1,
                            "dept": "Finance & Account-CC",
                            "particulars": "Cash receive from front office 3 April 26 and 4 April 26 morning",
                            "amount_kyats": 128037698.0,
                        },
                    ),
                    (
                        0,
                        0,
                        {
                            "line_type": "payment",
                            "sequence": 1,
                            "dept": "",
                            "particulars": "Commission Fee for Feb 2026 (Opal)",
                            "amount_kyats": 1516850.0,
                        },
                    ),
                ],
            }
        )

    def test_xlsx_layout_and_victoria_colours(self):
        content = self.register._generate_xlsx()
        self.assertTrue(content[:2] == b"PK")

        from openpyxl import load_workbook

        book = load_workbook(BytesIO(content))
        sheet = book["Daily Cash Register"]
        self.assertEqual(sheet["A2"].value, "Daily Cash Register")
        self.assertEqual(sheet["A3"].value, "No")
        self.assertEqual(sheet["C4"].value, "Opening Cash balance")
        self.assertEqual(sheet["D4"].value, 1029986000)
        self.assertEqual(sheet["C5"].value, "Cash receive from front office 3 April 26 and 4 April 26 morning")
        self.assertEqual(sheet["D5"].value, 128037698)
        self.assertEqual(sheet["C6"].value, "Total Cash Receipt")
        self.assertEqual(sheet["C7"].value, "Commission Fee for Feb 2026 (Opal)")
        self.assertEqual(sheet["C8"].value, "Total Cash Payment")
        self.assertEqual(sheet["C9"].value, "Closing Cash balance")

        self.assertEqual(_fill_rgb(sheet["A3"]), COLOR_HEADER.lstrip("#").upper())
        self.assertEqual(_fill_rgb(sheet["C4"]), COLOR_BALANCE.lstrip("#").upper())
        self.assertEqual(_fill_rgb(sheet["C6"]), COLOR_TOTAL.lstrip("#").upper())
        self.assertEqual(_fill_rgb(sheet["C8"]), COLOR_TOTAL.lstrip("#").upper())
        self.assertEqual(_fill_rgb(sheet["C9"]), COLOR_BALANCE.lstrip("#").upper())

    def test_reference_particulars_and_manual_override(self):
        reference = "INV/2026/09/21/0007"
        aml = SimpleNamespace(
            cash_register_particulars=False,
            move_id=SimpleNamespace(ref=reference),
            name="Manual Payment: " + reference,
        )
        self.assertEqual(self.register._particulars_from_aml(aml), reference)
        aml.move_id.ref = False
        self.assertEqual(self.register._particulars_from_aml(aml), aml.name)
        aml.cash_register_particulars = "Cash received from reception"
        self.assertEqual(self.register._particulars_from_aml(aml), aml.cash_register_particulars)

    def test_legacy_particulars_preserve_edits(self):
        reference = "INV/2026/09/21/0007"
        aml = SimpleNamespace(
            cash_register_particulars=False,
            move_id=SimpleNamespace(ref=reference),
            name="Manual Payment: " + reference,
        )
        line = SimpleNamespace(
            ensure_one=Mock(), move_line_id=aml, register_id=self.register,
            particulars=reference + " - " + aml.name,
        )
        method = type(self.env["vhg.daily.cash.register.line"])._report_particulars
        self.assertEqual(method(line), reference)
        line.particulars = "Cash collected by cashier"
        self.assertEqual(method(line), line.particulars)

    def test_denominations_and_signature_blocks(self):
        self.register.denom_line_ids = [
            (0, 0, {"code": "small", "name": "Small note", "sequence": 1}),
            (0, 0, {"code": "10000", "name": "10,000", "sequence": 2}),
        ]
        content = self.register._generate_xlsx()
        from openpyxl import load_workbook
        sheet = load_workbook(BytesIO(content))["Daily Cash Register"]
        values = [cell.value for row in sheet for cell in row]
        self.assertNotIn("Small note", values)
        denomination = next(row for row in sheet if row[1].value == "10,000")
        self.assertEqual(denomination[0].value, 1)
        with ZipFile(BytesIO(content)) as archive:
            drawing = etree.fromstring(archive.read("xl/drawings/drawing1.xml"))
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        shapes = drawing.findall("{*}twoCellAnchor/{*}sp")
        self.assertEqual(len(shapes), 3)
        self.assertEqual([
            shape.xpath(".//a:t/text()", namespaces=ns) for shape in shapes
        ], [
            ["Confirmed by;", "Other Department"],
            ["Checked By;", "DCA/ACA"],
            ["Prepared by:", "Chief Cashier/Asst; Chief Cashier"],
        ])
